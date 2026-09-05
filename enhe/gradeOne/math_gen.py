# -*- coding:utf-8 -*-
# ！/usr/bin/env python
# @Time : 2026/9/5 15:44
# @Author : waxberry
# @File : math_gen.py
# @Software : PyCharm



import random
import configparser
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml.shared import OxmlElement, qn as qn_


def load_config(config_path="config.ini"):
    cfg = configparser.ConfigParser()
    cfg.read(config_path, encoding="utf-8")

    exercise_conf = cfg["exercise"]
    font_conf = cfg["font"]
    output_conf = cfg["output"]

    columns = int(exercise_conf["columns"])
    per_col_q = int(exercise_conf["per_column_question"])

    conf = {
        "max_number": int(exercise_conf["max_number"]),
        "columns": columns,
        "per_column_question": per_col_q,
        "question_count": columns * per_col_q,
        "only_add": exercise_conf.getboolean("only_add"),
        "only_sub": exercise_conf.getboolean("only_sub"),
        "sub_no_negative": exercise_conf.getboolean("sub_no_negative"),
        "enable_answer": exercise_conf.getboolean("enable_answer"),

        "font_name": font_conf["font_name"],
        "title_font_size": int(font_conf["title_font_size"]),
        "content_font_size": int(font_conf["content_font_size"]),

        "file_name": output_conf["file_name"]
    }
    return conf


def generate_math_exercises(max_num: int, count: int, only_add: bool, only_sub: bool, sub_no_negative: bool):
    exercises = []
    for _ in range(count):
        if only_add:
            a = random.randint(0, max_num)
            b = random.randint(0, max_num - a)
            q = f"{a} + {b} ="
            ans = a + b
        elif only_sub:
            if sub_no_negative:
                a = random.randint(0, max_num)
                b = random.randint(0, a)
            else:
                a = random.randint(0, max_num)
                b = random.randint(0, max_num)
            q = f"{a} - {b} ="
            ans = a - b
        else:
            if random.choice([True, False]):
                a = random.randint(0, max_num)
                b = random.randint(0, max_num - a)
                q = f"{a} + {b} ="
                ans = a + b
            else:
                if sub_no_negative:
                    a = random.randint(0, max_num)
                    b = random.randint(0, a)
                else:
                    a = random.randint(0, max_num)
                    b = random.randint(0, max_num)
                q = f"{a} - {b} ="
                ans = a - b
        exercises.append((q, ans))
    return exercises


def set_font(run, font_name, font_size):
    run.font.size = Pt(font_size)
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)


def remove_cell_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for border_name in ["top", "left", "bottom", "right"]:
        tag = OxmlElement(f'w:{border_name}')
        tag.set(qn_('w:val'), "none")
        tcBorders.append(tag)
    tcPr.append(tcBorders)


def add_table_row(doc, item_list, col_num, font_name, font_size):
    row = doc.add_table(rows=1, cols=col_num)
    row.alignment = WD_TABLE_ALIGNMENT.LEFT
    for idx, cell in enumerate(row.rows[0].cells):
        remove_cell_border(cell)
        p = cell.paragraphs[0]
        run = p.add_run(item_list[idx])
        set_font(run, font_name, font_size)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def save_to_word_table(exercises, conf):
    doc = Document()
    # 标题
    title = doc.add_paragraph()
    run = title.add_run("数学练习题")
    set_font(run, conf["font_name"], conf["title_font_size"])
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("")

    col = conf["columns"]
    # 输出题目
    for i in range(0, len(exercises), col):
        batch = exercises[i:i + col]
        q_list = [x[0] for x in batch]
        while len(q_list) < col:
            q_list.append("")
        add_table_row(doc, q_list, col, conf["font_name"], conf["content_font_size"])

    # 是否输出参考答案
    if conf["enable_answer"]:
        doc.add_paragraph("")
        ans_title = doc.add_paragraph()
        ans_run = ans_title.add_run("参考答案")
        set_font(ans_run, conf["font_name"], conf["title_font_size"])
        doc.add_paragraph("")
        for i in range(0, len(exercises), col):
            batch = exercises[i:i + col]
            ans_list = [f"{q}{ans}" for q, ans in batch]
            while len(ans_list) < col:
                ans_list.append("")
            add_table_row(doc, ans_list, col, conf["font_name"], conf["content_font_size"])

    doc.save(conf["file_name"])
    print(f"✅ Word文档生成成功：{conf['file_name']}")
    print(f"📝 配置：每行{conf['columns']}列，每列{conf['per_column_question']}道，总题数{conf['question_count']}")


if __name__ == "__main__":
    config = load_config("config.ini")
    exercise_list = generate_math_exercises(
        max_num=config["max_number"],
        count=config["question_count"],
        only_add=config["only_add"],
        only_sub=config["only_sub"],
        sub_no_negative=config["sub_no_negative"]
    )
    save_to_word_table(exercise_list, config)
